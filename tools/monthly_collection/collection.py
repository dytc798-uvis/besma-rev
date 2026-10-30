"""Daily NAVERWORKS intake; immutable source receipts and conservative workbook updates."""
from __future__ import annotations
import argparse,collections,contextlib,datetime as dt,hashlib,json,os,re,shutil,sqlite3,sys,uuid,zipfile,unicodedata
from pathlib import Path
KST=dt.timezone(dt.timedelta(hours=9))
BASE=Path(r'D:\JSI\safety-workbench')
STATE=BASE/'data/monthly-collection'
NAS=Path(r'Z:\4. 안전보건관리실\★월간 자료 취합')
NAS_LINK_ROOT=NAS
PAYLOAD=Path(r'E:\JSI\backups\monthly-collection')
CATEGORIES={'theme':'1. 안전 점검/테마점검','legal-training':'2. 법적교육(특별,정기 등)','risk':'3. 위험성평가','opinions':'4. 근로자 의견 청취','safety-cost':'6. 산업안전보건관리비'}
LABELS={'theme':'테마점검','legal-training':'법정교육','risk':'위험성평가','opinions':'근로자 의견청취','safety-cost':'산업안전보건관리비'}
ALLOWED={'.zip','.xlsx','.xlsm','.xls','.hwp','.hwpx','.docx','.doc','.pptx','.pdf','.png','.jpg','.jpeg','.txt','.csv'}
def now():return dt.datetime.now(KST)
def read(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))

def payload(root):
 config=root/'config.json'
 value=read(config).get('payload_root') if config.exists() else None
 if not value:return root
 if root.resolve()!=STATE.resolve() or Path(value).resolve()!=PAYLOAD.resolve():raise ValueError('PAYLOAD_ROOT_NOT_APPROVED')
 return PAYLOAD
def atomic(p,v):
 p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);tmp=p.with_name(p.name+'.'+uuid.uuid4().hex+'.tmp');tmp.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8');os.replace(tmp,p)
def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def key(*parts):return hashlib.sha256(json.dumps(parts,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
def norm(s):return re.sub(r'[^a-z0-9가-힣]','',str(s).lower())
def db(root):
 root.mkdir(parents=True,exist_ok=True);c=sqlite3.connect(root/'ledger.sqlite3',timeout=15);c.row_factory=sqlite3.Row
 c.executescript('''PRAGMA journal_mode=WAL;
 CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY,category TEXT,month TEXT,team TEXT,source_path TEXT,sha256 TEXT,size INTEGER,first_seen TEXT,last_seen TEXT,nas_path TEXT,site TEXT,site_code TEXT,state TEXT,UNIQUE(category,month,team,source_path,sha256));
 CREATE TABLE IF NOT EXISTS opinions(id TEXT PRIMARY KEY,source_id TEXT,site TEXT,person TEXT,day TEXT,opinion TEXT,action TEXT,status TEXT,source TEXT,month TEXT,team TEXT,updated_at TEXT);
 CREATE TABLE IF NOT EXISTS opinion_sources(origin TEXT PRIMARY KEY,opinion_id TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS reviews(id TEXT PRIMARY KEY,kind TEXT,source_id TEXT,reason TEXT,created_at TEXT,resolved INTEGER DEFAULT 0);
 CREATE TABLE IF NOT EXISTS cells(book TEXT,sheet TEXT,address TEXT,value TEXT,PRIMARY KEY(book,sheet,address));
 CREATE TABLE IF NOT EXISTS publications(id TEXT PRIMARY KEY,at TEXT,state TEXT,details TEXT);
 CREATE TABLE IF NOT EXISTS deliveries(file_id TEXT PRIMARY KEY,source_sha256 TEXT,source_size INTEGER,sha256 TEXT,size INTEGER,recipe TEXT,state TEXT,proof TEXT,at TEXT);
 ''');return c
@contextlib.contextmanager
def locked(root,lock_name='worker.lock'):
 root.mkdir(parents=True,exist_ok=True)
 with (root/lock_name).open('a+b') as f:
  if os.name=='nt':
   import msvcrt
   f.write(b'0');f.flush();f.seek(0);msvcrt.locking(f.fileno(),msvcrt.LK_NBLCK,1)
  try:yield
  finally:
   if os.name=='nt':f.seek(0);msvcrt.locking(f.fileno(),msvcrt.LK_UNLCK,1)
def request(root,manual=False,clock=None):
 cfg=read(root/'config.json');stamp=clock or now();due=stamp.replace(hour=2,minute=0,second=0,microsecond=0)
 (root/'requests').mkdir(exist_ok=True)
 if manual:
  for prior in sorted((root/'requests').glob('*.json')):
   if not prior.name.endswith('.done.json') and not Path(str(prior)+'.done.json').exists():return {'state':'QUEUED','path':str(prior)}
 if stamp<due:due-=dt.timedelta(days=1)
 if not manual and (not cfg['enabled'] or due<dt.datetime.fromisoformat(cfg['start_at'])):return {'state':'NOT_DUE'}
 date=stamp.date() if manual else due.date();months=[(date.replace(day=1)-dt.timedelta(days=1)).strftime('%Y-%m'),date.strftime('%Y-%m')]
 following=(date.replace(day=28)+dt.timedelta(days=4)).replace(day=1).strftime('%Y-%m');months.append(following)
 target=root/'requests'/(date.isoformat()+('-manual-'+stamp.strftime('%H%M%S%f') if manual else '-0200')+'.json')
 if target.exists():return {'state':'QUEUED','path':str(target)}
 value={'id':uuid.uuid4().hex,'run_date':date.isoformat(),'requested_at':stamp.isoformat(),'months':months,'main_month':date.strftime('%Y-%m'),'risk_month':following}
 atomic(target,value);return {'state':'QUEUED','path':str(target)}
def catalog(root):
 # Only a verified shared ERP receipt is used; never log credential/pay columns.
 sites={};receipt_dir=BASE/'data/erp-collection/latest/site_master'
 receipts=sorted(receipt_dir.glob('*.json'),key=lambda p:p.stat().st_mtime,reverse=True)
 if receipts:
  j=read(receipts[0]);p=Path(j['path'])
  if j['state']=='RECEIVED' and p.is_file() and p.stat().st_size==j['size'] and sha(p)==j['raw_sha256']:
   for x in j.get('safe_rows',[]):sites[x['code']]={'code':x['code'],'site':x['site']}
 server=root/'site-catalog.json'
 if server.exists():
  for row in read(server)['sites']:
   code=row['code'];prior=sites.get(code);item=dict(row);item['aliases']=[prior['site']] if prior else []
   team=re.match(r'^\[(?:준공-)?([1-6])\.',item['site']);item['team']=team[1]+'팀' if team else '미분류';sites[code]=item
 return list(sites.values())
def identify(text,sites):
 from site_binding import bind
 decision=bind(text,sites)
 if decision['state'].startswith('EXACT'):return next(site for site in sites if site['code']==decision['code'])
 if decision['state']=='AMBIGUOUS':return None
 codes=[s for s in sites if re.search(r'(?<!\d)'+re.escape(s['code'])+r'(?!\d)',text)]
 if len(codes)==1:return codes[0]
 value=norm(text);hits=[]
 for s in sites:
  labels=[norm(re.sub(r'^\[[^]]+\]','',v)) for v in [s['site'],*s.get('aliases',[])]]
  if any(len(label)>=8 and label in value for label in labels):hits.append(s)
 return hits[0] if len(hits)==1 else None
def safe_parts(name):
 parts=name.replace('\\','/').split('/')
 if any(not p or p in ('.','..') or re.search(r'[\x00-\x1f:*?"<>|]',p) or p.rstrip(' .')!=p or re.fullmatch(r'(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\..*)?',p) for p in parts):raise ValueError('UNSAFE_ARCHIVE_PATH')
 return parts
def below(path,root):
 if not path.resolve().is_relative_to(root.resolve()):raise ValueError('DESTINATION_ESCAPE')
 p=path
 while p!=root:
  if p.exists() and (p.is_symlink() or getattr(p.stat(),'st_file_attributes',0)&0x400):raise ValueError('REPARSE_DESTINATION')
  p=p.parent
def verified_copy(source,target):
 before=source.stat();digest=sha(source)
 # RaiDrive/WebDAV may return ACCESS_DENIED rather than ENOENT for a missing
 # grandparent. Create each missing level from the existing parent downward.
 missing=[];cursor=target.parent
 while not cursor.exists():
  missing.append(cursor)
  if cursor==cursor.parent:raise FileNotFoundError('DESTINATION_ROOT_MISSING')
  cursor=cursor.parent
 for directory in reversed(missing):directory.mkdir(exist_ok=True)
 if target.exists():
  if sha(target)==digest:return digest
  raise FileExistsError('CONTENT_COLLISION')
 temp=target.with_name(target.name+'.'+uuid.uuid4().hex+'.partial')
 try:
  with source.open('rb') as src,temp.open('xb') as out:shutil.copyfileobj(src,out);out.flush();os.fsync(out.fileno())
  if (before.st_size,before.st_mtime_ns)!=(source.stat().st_size,source.stat().st_mtime_ns) or sha(source)!=digest or sha(temp)!=digest:raise ValueError('UNSTABLE_SOURCE')
  # Link/rename must never silently overwrite another writer's file.
  if target.exists():raise FileExistsError('DESTINATION_CHANGED')
  temp.rename(target)
  if sha(target)!=digest:raise ValueError('SAVED_HASH_MISMATCH')
  return digest
 finally:
  if temp.exists():temp.unlink()
def texts(path):
 ext=path.suffix.lower()
 if ext in ('.xlsx','.xlsm'):
  import openpyxl
  w=openpyxl.load_workbook(path,read_only=True,data_only=True);result=[]
  try:
   for s in w:
    for row in s.iter_rows(max_row=min(s.max_row,2000),max_col=min(s.max_column,40)):
     result.append(' '.join(str(c.value or '') for c in row))
  finally:w.close()
  return '\n'.join(result)
 if ext=='.pdf':
  import fitz
  with fitz.open(path) as d:
   if len(d)>200:raise ValueError('PDF_PAGE_LIMIT')
   return '\n'.join(p.get_text() for p in d)[:500000]
 if ext in ('.txt','.csv'):return path.read_text(encoding='utf-8-sig',errors='replace')[:500000]
 return ''
def review(c,kind,source,reason):
 c.execute('INSERT OR IGNORE INTO reviews(id,kind,source_id,reason,created_at) VALUES(?,?,?,?,?)',(key(kind,source,reason),kind,source,reason,now().isoformat()))

def optimization_enabled(root):
 return (root/'config.json').exists() and read(root/'config.json').get('image_optimization',{}).get('enabled') is True

def deliver(c,root,identity,source,dest,old=None):
 """Original hash stays in files; actual NAS hash and size have a separate receipt."""
 original_sha=sha(source);proof={'state':'DISABLED','recipe':'none','source_sha256':original_sha,'source_size':source.stat().st_size,'sha256':original_sha,'size':source.stat().st_size};delivery=source
 if optimization_enabled(root):
  from image_optimizer import optimize
  delivery,proof=optimize(source,payload(root))
 prior=c.execute('SELECT * FROM deliveries WHERE file_id=?',(identity,)).fetchone()
 expected=prior['sha256'] if prior else (old['sha256'] if old else None)
 if dest.exists() and sha(dest)!=proof['sha256']:
  if old and str(dest)==old['nas_path'] and sha(dest)==expected:
   stage=dest.with_name(dest.name+'.'+uuid.uuid4().hex+'.optimizing');verified_copy(delivery,stage)
   try:
    if sha(dest)!=expected:raise ValueError('NAS_CHANGED_DURING_OPTIMIZATION')
    os.replace(stage,dest)
    if sha(dest)!=proof['sha256']:raise ValueError('OPTIMIZED_DELIVERY_HASH')
   finally:
    if stage.exists():stage.unlink()
  else:
   dest=dest.with_name(dest.stem+'_'+proof['sha256'][:12]+dest.suffix);verified_copy(delivery,dest)
 else:verified_copy(delivery,dest)
 c.execute('INSERT OR REPLACE INTO deliveries VALUES(?,?,?,?,?,?,?,?,?)',(identity,original_sha,source.stat().st_size,proof['sha256'],proof['size'],proof['recipe'],proof['state'],json.dumps(proof,ensure_ascii=False),now().isoformat()))
 if proof['state']=='PRESERVED_CHECK_REQUIRED':review(c,'IMAGE_OPTIMIZATION',identity,'그림 축소 검증 보류 · 원본 보존')
 else:c.execute("UPDATE reviews SET resolved=1 WHERE kind='IMAGE_OPTIMIZATION' AND source_id=? AND reason='그림 축소 검증 보류 · 원본 보존'",(identity,))
 return dest,proof
def day_value(value):
 if isinstance(value,(dt.datetime,dt.date)):return value.strftime('%Y-%m-%d')
 s=str(value or '').strip()
 for fmt in ('%Y-%m-%d','%Y.%m.%d','%y.%m.%d','%Y/%m/%d'):
  try:return dt.datetime.strptime(s,fmt).strftime('%Y-%m-%d')
  except ValueError:pass
 return ''
def opinion_rows(path,month):
 if path.suffix.lower() not in ('.xlsx','.xlsm'):return []
 import openpyxl
 w=openpyxl.load_workbook(path,read_only=True,data_only=True);out=[]
 try:
  for s in w:
   at=lambda a:str(s[a].value or '').replace(' ','')
   a='의견일자' in at('C4') and '내용' in at('G5');b='일자' in at('C4') and '내용' in at('H5') and '조치' in at('I5')
   if not(a or b) or '현장명' not in at('D4'):continue
   for row_number,row in enumerate(s.iter_rows(min_row=6,max_row=min(s.max_row,2000),max_col=13,values_only=True),6):
    day=day_value(row[2]);site=str(row[3] or '').strip();person=str(row[4] or '').strip();opinion=str(row[6 if a else 7] or '').strip();action=str(row[10 if a else 9] or '').strip()
    if not any((site,person,opinion,action)):continue
    if not day or not site or not opinion or day[:7]!=month:continue
    none=bool(re.fullmatch(r'(해당사항없음|해당없음|의견없음|없음)[.。]?',norm(opinion)))
    out.append({'site':site,'person':person,'day':day,'opinion':opinion,'action':action,'status':'의견 없음' if none else '완료 미확인','source_row':s.title+'!'+str(row_number)})
 finally:w.close()
 return out
def ingest(root,manifest,nas=NAS,sites=None):
 c=db(root)
 try:
  result=_ingest(c,root,manifest,nas,sites)
  if nas==NAS:
   snapshot(root,c)
   try:
    from publish_server import publish as publish_metadata
    result['server']=publish_metadata()
   except Exception as error:
    result['server']={'state':'NOT_PUBLISHED','error':type(error).__name__};atomic(root/'failures'/('metadata-'+now().strftime('%Y%m%d_%H%M%S%f')+'.json'),result['server'])
   if result.get('state')=='NAS_SAVED':
    # A later profile/download can fail. Already receipted originals should
    # still enter the existing search refresh, independently of workbook finish.
    try:
     from search_refresh import request as request_search_refresh
     result['search_refresh']=request_search_refresh([],root,nas)
    except Exception as error:
     result['search_refresh']={'state':'RETRY_PENDING','error':type(error).__name__,'received_files_preserved':True}
  return result
 finally:c.close()
def _ingest(c,root,manifest,nas,sites):
 j=read(manifest);category=j['category'];month=j['month'];dt.datetime.strptime(month,'%Y-%m');folder=CATEGORIES[category]
 if not nas.is_dir():raise FileNotFoundError('NAS_NOT_AVAILABLE')
 sites=catalog(root) if sites is None else sites;counts=collections.Counter();receipts=[]
 for archive in j['archives']:
  team=archive.get('Folder',archive.get('folder'));zp=archive.get('ZipPath',archive.get('zip_path'))
  if not re.fullmatch('[1-6]팀|관급|미분류',team):raise ValueError('TEAM_NOT_ALLOWED')
  if not zp:counts['empty_teams']+=1;continue
  storage=payload(root);zip_path=Path(zp);below(zip_path,storage/'downloads');frozen=storage/'archives'/sha(zip_path)/'source.zip';verified_copy(zip_path,frozen)
  with zipfile.ZipFile(frozen) as z:
   infos=z.infolist()
   if len(infos)>5000 or sum(x.file_size for x in infos)>2*1024**3:raise ValueError('ARCHIVE_LIMIT')
   names=set()
   for info in infos:
    if info.is_dir():continue
    parts=safe_parts(info.filename);canonical=unicodedata.normalize('NFC','/'.join(parts)).casefold()
    if canonical in names or (info.external_attr>>16)&0xF000==0xA000:raise ValueError('DUPLICATE_OR_LINK_ENTRY')
    names.add(canonical)
    if info.file_size>500*1024**2:raise ValueError('FILE_SIZE_LIMIT')
    if Path(parts[-1]).suffix.lower() not in ALLOWED or parts[-1].lower() in ('thumbs.db','desktop.ini') or parts[-1].startswith('~$') or any(p.casefold() in ('thumbs','__macosx') for p in parts):continue
    if re.search(r'(?i)(password|api[_ -]?key|access[_ -]?token|비밀번호|개인키)',parts[-1]):counts['excluded_credentials']+=1;continue
    if not info.file_size:counts['zero_byte']+=1;continue
    data=z.read(info);digest=hashlib.sha256(data).hexdigest();source=storage/'sources'/digest/parts[-1];source.parent.mkdir(parents=True,exist_ok=True)
    if not source.exists():
     with source.open('xb') as f:f.write(data)
    if sha(source)!=digest:raise ValueError('SOURCE_HASH_CHANGED')
    source_path='/'.join(parts);identity=key(category,month,team,source_path,digest)
    old=c.execute('SELECT * FROM files WHERE id=?',(identity,)).fetchone()
    prior_delivery=c.execute('SELECT * FROM deliveries WHERE file_id=?',(identity,)).fetchone()
    expected=prior_delivery['sha256'] if prior_delivery else digest
    # An unchanged original with the same recipe does not need another failed candidate every night.
    # Explicit optimize-existing --retry-checks remains available for a reviewed retry.
    optimized_current=not optimization_enabled(root) or (prior_delivery is not None and prior_delivery['recipe']==__import__('image_optimizer').RECIPE)
    if old and Path(old['nas_path']).is_file() and sha(Path(old['nas_path']))==expected and optimized_current:
     _rebind_file(c,old,sites)
     c.execute('UPDATE files SET last_seen=? WHERE id=?',(now().isoformat(),identity));counts['unchanged']+=1;continue
    extracted=''
    try:extracted=texts(source)
    except Exception:review(c,'EXTRACTION',identity,'원문 확인 필요')
    site=identify(source_path+'\n'+extracted,sites)
    mapped_team=site.get('team',team) if team=='미분류' and site else team
    dest=nas/folder/(month[2:].replace('-','.')+'월')/mapped_team
    # Original folder names are preserved. Unknown site filenames remain in the team folder for review.
    dest=dest.joinpath(*parts);below(dest,nas)
    dest,delivery_proof=deliver(c,root,identity,source,dest,old);below(dest,nas)
    stamp=now().isoformat();state='SAVED' if site else 'SITE_REVIEW'
    prior=c.execute('SELECT first_seen FROM files WHERE category=? AND month=? AND team=? AND source_path=? ORDER BY first_seen LIMIT 1',(category,month,team,source_path)).fetchone()
    c.execute('INSERT OR REPLACE INTO files VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(identity,category,month,mapped_team,source_path,digest,len(data),prior['first_seen'] if prior else stamp,stamp,str(dest),site['site'] if site else '',site['code'] if site else '',state))
    if site is None:review(c,'SITE',identity,'현장 일치 확인 필요')
    if category=='opinions':
     rows=opinion_rows(source,month)
     if not rows:review(c,'OPINION',identity,'의견 원문·날짜·현장 확인 필요')
     for record in rows:
      oid=key(norm(record['site']),record['day'],norm(record['person']),norm(record['opinion']))
      origin=key(category,month,team,source_path,record['source_row']);binding=c.execute('SELECT opinion_id FROM opinion_sources WHERE origin=?',(origin,)).fetchone()
      if binding:
       old=c.execute('SELECT * FROM opinions WHERE id=?',(binding['opinion_id'],)).fetchone()
       if old and (norm(old['site'])!=norm(record['site']) or old['day']!=record['day'] or norm(old['person'])!=norm(record['person'])):
        review(c,'OPINION_IDENTITY',identity,'수정본의 현장·의견일·제안자가 달라 원문 확인 필요');continue
       if old:oid=old['id']
      existing=c.execute('SELECT * FROM opinions WHERE id=?',(oid,)).fetchone()
      status=existing['status'] if existing and existing['status']=='완료' else record['status']
      action=record['action'] or (existing['action'] if existing else '')
      if existing and existing['action'] and action and action!=existing['action'] and action not in existing['action']:action=existing['action']+'\n'+action
      c.execute('INSERT OR REPLACE INTO opinions VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(oid,identity,record['site'],record['person'],record['day'],record['opinion'],action,status,str(dest),month,mapped_team,stamp))
      c.execute('INSERT OR REPLACE INTO opinion_sources VALUES(?,?)',(origin,oid))
    counts['saved']+=1;receipts.append({'id':identity,'source_sha256':digest,'source_size':len(data),'sha256':delivery_proof['sha256'],'size':delivery_proof['size'],'optimization':delivery_proof['state'],'nas_path':str(dest)})
  c.commit()
 result={'at':now().isoformat(),'state':'NAS_SAVED','category':category,'month':month,'counts':dict(counts),'files':receipts};atomic(root/'receipts'/(key(str(manifest))+'.json'),result);return result
def value_json(value):
 if isinstance(value,float) and value.is_integer():value=int(value)
 return json.dumps(value,ensure_ascii=False,default=str)
def native_copy(source,*arguments):
 # Excel's late-bound Copy property may execute before Python applies kwargs.
 # Invoke the public IDispatch method with explicitly typed missing arguments.
 import win32com.client,pythoncom
 missing=win32com.client.VARIANT(pythoncom.VT_ERROR,-2147352572)
 values=[missing if x is None else x._oleobj_ for x in arguments]
 source._oleobj_.Invoke(source._oleobj_.GetIDsOfNames('Copy'),0,pythoncom.DISPATCH_METHOD,False,*values)
def set_managed(c,book,sheet,address,value,bookkey,changes):
 cell=sheet.Range(address);current=cell.Value2;last=c.execute('SELECT value FROM cells WHERE book=? AND sheet=? AND address=?',(bookkey,sheet.Name,address)).fetchone()
 wanted=value_json(value)
 if last and value_json(current)!=last['value'] and value_json(current)!=wanted:
  review(c,'MANUAL_CELL',key(bookkey,sheet.Name,address),'사용자 수정 셀 보존');return False
 if not last and current not in (None,'') and value_json(current)!=wanted:
  review(c,'EXISTING_CELL',key(bookkey,sheet.Name,address),'기존 입력값 보존');return False
 if value_json(current)!=wanted:
  if isinstance(value,str):cell.NumberFormat='@'
  cell.Value2=value;changes.append((sheet.Name,address))
 c.execute('INSERT OR REPLACE INTO cells VALUES(?,?,?,?)',(bookkey,sheet.Name,address,wanted));return True
def workbook_transaction(root,c,target,template,mutate):
 import win32com.client,pythoncom
 run=payload(root)/'publications'/now().strftime('%Y%m%d_%H%M%S_%f');run.mkdir(parents=True)
 baseline=target if target.exists() else template;before=sha(baseline) if baseline else None
 candidate=run/target.name
 if baseline:shutil.copy2(baseline,candidate);shutil.copy2(baseline,run/'baseline.xlsx')
 excel=None;book=None
 try:
  excel=win32com.client.DispatchEx('Excel.Application');excel.Visible=False;excel.DisplayAlerts=False;excel.AutomationSecurity=3
  book=excel.Workbooks.Open(str(candidate),UpdateLinks=0) if baseline else excel.Workbooks.Add()
  changed=mutate(book,not target.exists())
  if not changed and target.exists():
   book.Close(False);book=None
   if sha(target)!=before:raise ValueError('WORKBOOK_CHANGED_DURING_BUILD')
   # These are this transaction's identical temporary copies, never source/user files.
   for temporary in (candidate,run/'baseline.xlsx'):
    if sha(temporary)!=before:raise ValueError('UNCHANGED_WORKBOOK_COPY_CHANGED')
    temporary.unlink()
   receipt={'target':str(target),'sha256':before,'baseline_sha256':before,'changed_cells':0,'at':now().isoformat(),'image_optimization':{'state':'UNCHANGED'},'state':'UNCHANGED'}
   atomic(run/'receipt.json',receipt);return receipt
  excel.CalculateFullRebuild()
  if baseline:book.Save()
  else:book.SaveAs(str(candidate),FileFormat=51)
  # Native reopen verifies the bytes that will be published.
  book.Close(False);book=excel.Workbooks.Open(str(candidate),UpdateLinks=0,ReadOnly=True)
  for name,address in changed:
   cell=book.Worksheets(name).Range(address)
   if str(cell.Text).startswith(('#REF!','#VALUE!','#DIV/0!')):raise ValueError('WRITTEN_CELL_ERROR')
  # One actual affected page is retained on D for layout checks, never on the NAS.
  if changed:
   book.Worksheets(changed[0][0]).ExportAsFixedFormat(0,str(run/'preview.pdf'),From=1,To=1)
  book.Close(False);book=None
  delivery=candidate;image_proof={'state':'DISABLED'}
  if optimization_enabled(root):
   from image_optimizer import optimize
   delivery,image_proof=optimize(candidate,payload(root))
   if image_proof['state']=='PRESERVED_CHECK_REQUIRED':review(c,'WORKBOOK_IMAGE',key(str(target)),'대장 그림 축소 검증 보류 · 원본 품질 유지')
  if baseline and sha(baseline)!=before:raise ValueError('WORKBOOK_CHANGED_DURING_BUILD')
  target.parent.mkdir(parents=True,exist_ok=True);staging=target.with_name(target.name+'.'+uuid.uuid4().hex+'.publishing');verified_copy(delivery,staging)
  if target.exists() and sha(target)!=before:raise ValueError('TARGET_CHANGED_BEFORE_COMMIT')
  os.replace(staging,target)
  if sha(target)!=sha(delivery):raise ValueError('WORKBOOK_PUBLISH_HASH')
  receipt={'target':str(target),'sha256':sha(target),'baseline_sha256':before,'changed_cells':len(changed),'at':now().isoformat(),'image_optimization':image_proof};atomic(run/'receipt.json',receipt);return receipt
 finally:
  if book is not None:book.Close(False)
  if excel is not None:excel.Quit()
def prepare_evaluation(root,c,nas):
 target=nas/'월간자료 취합현황.xlsx'
 rows=c.execute('SELECT * FROM files ORDER BY month,team,site,category,source_path,first_seen').fetchall()
 def mutate(book,fresh):
  changes=[]
  if fresh:
   while book.Worksheets.Count>1:book.Worksheets(book.Worksheets.Count).Delete()
   book.Worksheets(1).Name='취합현황';book.Worksheets.Add(After=book.Worksheets(1)).Name='종합평가 반영자료'
  sheet=book.Worksheets('취합현황');evaluation=book.Worksheets('종합평가 반영자료')
  headers=['자료월','팀','현장명','자료종류','파일명','최초 확인일시','저장상태','증빙 파일']
  for sh in [sheet,evaluation]:
   if fresh:
    sh.Range('A1:H1').Value=(tuple(headers),);sh.Range('A1:H1').Font.Bold=True;sh.Range('A1:H1').Interior.Color=0xEEEEEE;sh.Columns('A:H').ColumnWidth=18;sh.Columns('C').ColumnWidth=45;sh.Columns('E').ColumnWidth=48;sh.Columns('H').ColumnWidth=55;sh.PageSetup.Orientation=2;sh.PageSetup.Zoom=False;sh.PageSetup.FitToPagesWide=1;sh.PageSetup.FitToPagesTall=False;sh.PageSetup.PrintTitleRows='$1:$1'
  # Row mapping stays in D; manual row edits do not move an input to another site.
  for item in rows:
   for sh in ([sheet,evaluation] if item['site'] else [sheet]):
    rowkey='row:'+item['id'];binding=c.execute('SELECT value FROM cells WHERE book=? AND sheet=? AND address=?',(str(target),sh.Name,rowkey)).fetchone()
    if binding:r=int(json.loads(binding['value']))
    else:
     r=max(2,sh.Cells(sh.Rows.Count,1).End(-4162).Row+1);c.execute('INSERT INTO cells VALUES(?,?,?,?)',(str(target),sh.Name,rowkey,str(r)))
    vals=[item['month'],item['team'],item['site'] or '현장 확인 필요',LABELS[item['category']],Path(item['source_path']).name,item['first_seen'][:19].replace('T',' '),'저장 완료','증빙 파일 열기']
    for n,v in enumerate(vals):set_managed(c,book,sh,f'{chr(65+n)}{r}',v,str(target),changes)
    cell=sh.Range(f'H{r}');link_root=NAS_LINK_ROOT if Path(item['nas_path']).is_relative_to(NAS_LINK_ROOT) else nas;address=portable_link(item['nas_path'],link_root)
    if cell.HasFormula:review(c,'MANUAL_LINK',key(str(target),sh.Name,f'H{r}'),'사용자 수정 증빙 연결 보존');continue
    if cell.Hyperlinks.Count==0:
     sh.Hyperlinks.Add(Anchor=cell,Address=address,TextToDisplay=str(cell.Value2 or '증빙 파일 열기'));changes.append((sh.Name,f'H{r}'))
    elif str(cell.Hyperlinks.Item(1).Address)!=address:
     if managed_link(str(cell.Hyperlinks.Item(1).Address),item['nas_path'],link_root):
      replace_link(cell,sh,address);changes.append((sh.Name,f'H{r}'))
     else:review(c,'MANUAL_LINK',key(str(target),sh.Name,f'H{r}'),'사용자 수정 증빙 연결 보존')
  return changes
 return workbook_transaction(root,c,target,None,mutate)
def prepare_opinions(root,c,nas):
 rows=c.execute('SELECT * FROM opinions ORDER BY month,day,team,site,id').fetchall()
 for item in rows:
  if not item['month'].startswith('2026-'):review(c,'YEAR_TEMPLATE',item['id'],'해당 연도의 의견청취 대장 양식 확인 필요')
 rows=[item for item in rows if item['month'].startswith('2026-')]
 if not rows:return {'state':'NO_PARSED_OPINIONS','review_required':c.execute("SELECT COUNT(*) FROM reviews WHERE kind='OPINION' AND resolved=0").fetchone()[0]}
 folder=nas/CATEGORIES['opinions'];target=folder/'근로자 의견청취 대장_2026.xlsx';template=folder/'근로자 의견청취 대장_7월·8월 9월.xlsx'
 if not target.exists() and not template.is_file():raise FileNotFoundError('OPINION_TEMPLATE_MISSING')
 def mutate(book,fresh):
  changes=[];modified_months={}
  def existing_day(value):
   if isinstance(value,(int,float)) and 1<value<100000:return (dt.datetime(1899,12,30)+dt.timedelta(days=value+(1462 if book.Date1904 else 0))).strftime('%Y-%m-%d')
   return day_value(value)
  for item in rows:
   title=str(int(item['month'][5:]))+'월';names=[s.Name for s in book.Worksheets]
   if title not in names:
    source=book.Worksheets('9월');before_count=book.Worksheets.Count;native_copy(source,None,book.Worksheets(before_count))
    if book.Worksheets.Count!=before_count+1:raise ValueError('NATIVE_SHEET_COPY_FAILED')
    sh=book.Worksheets(book.Worksheets.Count);sh.Name=title
    footer=next((r for r in range(6,2500) if str(sh.Cells(r,2).Value2 or '').startswith('등록 ')),None)
    if not footer:raise ValueError('OPINION_LAYOUT_CHANGED')
    for i in range(sh.Shapes.Count,0,-1):
     shape=sh.Shapes(i)
     if 6<=shape.TopLeftCell.Row<footer:shape.Delete()
    sh.Range(f'B6:M{footer-1}').ClearContents();sh.Range('B2').Value2=f"{item['month'][:4]}년 {title} 근로자 의견청취 관리대장";sh.Range('B3').MergeArea.ClearContents()
   sh=book.Worksheets(title);footer=next((r for r in range(6,2500) if str(sh.Cells(r,2).Value2 or '').startswith('등록 ')),None)
   if not footer:raise ValueError('OPINION_FOOTER_MISSING')
   rowkey='row:'+item['id'];binding=c.execute('SELECT value FROM cells WHERE book=? AND sheet=? AND address=?',(str(target),title,rowkey)).fetchone()
   if binding:r=int(json.loads(binding['value']))
   else:
    # Exact business identity prevents importing a previously entered opinion again.
    same_day=[r for r in range(6,footer) if norm(sh.Cells(r,5).Value2)==norm(item['site']) and existing_day(sh.Cells(r,8).Value2)==item['day'] and (not item['person'] or norm(item['person']) in norm(sh.Cells(r,3).Value2))]
    match=[r for r in same_day if norm(sh.Cells(r,6).Value2)==norm(item['opinion'])]
    if len(match)==1:r=match[0]
    elif same_day:
     review(c,'EXISTING_OPINION',item['id'],'기존 대장의 같은 현장·날짜 의견과 원문 대조 필요');continue
    else:
     r=next((r for r in range(6,footer) if not sh.Cells(r,2).Value2),footer)
     if r==footer:sh.Rows(footer).Insert();native_copy(sh.Rows(footer-1),sh.Rows(footer));sh.Range(f'B{footer}:M{footer}').ClearContents();footer+=1
     modified_months[title]=footer-1
    c.execute('INSERT INTO cells VALUES(?,?,?,?)',(str(target),title,rowkey,str(r)))
   vals={2:r-5,3:item['team']+('\n'+item['person'] if item['person'] else ''),4:f"{item['month'][5:]}-{r-5:03}",5:item['site'],6:item['opinion'],8:item['day'],9:item['status'],11:item['action']+'\n증빙: '+Path(item['source']).name}
   for col,value in vals.items():set_managed(c,book,sh,f'{chr(64+col)}{r}',value,str(target),changes)
   if title in modified_months:
    last=footer-1;sh.Cells(footer,2).Formula=f'="등록 "&COUNT(B6:B{last})&"행 / 완료 "&COUNTIF(I6:I{last},"완료")&"건 / 진행중 "&COUNTIF(I6:I{last},"진행중")&"건 / 완료 미확인 "&COUNTIF(I6:I{last},"완료 미확인")&"건 / 의견 없음·공란 "&(COUNTIF(I6:I{last},"의견 없음")+COUNTIF(I6:I{last},"의견란 공란"))&"건"';sh.PageSetup.PrintArea=f'$B$1:$M${footer+1}'
  if '갑지' in [s.Name for s in book.Worksheets]:
   cover=book.Worksheets('갑지')
   for title,last in modified_months.items():
    total=next((r for r in range(11,60) if str(cover.Cells(r,2).Value2 or '')=='합계'),None)
    if total is None:raise ValueError('COVER_TOTAL_NOT_FOUND')
    positions=[r for r in range(11,total) if cover.Cells(r,2).Value2==title]
    if positions:r=positions[0]
    else:
     cover.Rows(total).Insert();native_copy(cover.Rows(total-1),cover.Rows(total));r=total;total+=1
    cover.Cells(r,2).Value2=title
    for col,label in [('F','완료'),('H','진행중'),('I','완료 미확인')]:cover.Range(f'{col}{r}').Formula=f'=COUNTIF(\'{title}\'!$I$6:$I${last},"{label}")'
    cover.Range(f'J{r}').Formula=f'=COUNTIF(\'{title}\'!$I$6:$I${last},"의견 없음")+COUNTIF(\'{title}\'!$I$6:$I${last},"의견란 공란")';cover.Range(f'L{r}').Formula=f'=COUNT(\'{title}\'!$B$6:$B${last})';cover.Range(f'D{r}').Formula=f'=L{r}-J{r}';cover.Range(f'K{r}').Formula=f'=IF(D{r}=0,0,F{r}/D{r})'
    for col in ['D','F','H','I','J','L']:cover.Range(f'{col}{total}').Formula=f'=SUM({col}11:{col}{total-1})'
    cover.Range(f'K{total}').Formula=f'=IF(D{total}=0,0,F{total}/D{total})';cover.Range('B7').Formula=f'=D{total}';cover.Range('E7').Formula=f'=F{total}';cover.Range('H7').Formula=f'=IF(D{total}=0,0,F{total}/D{total})';changes.append(('갑지',f'B{r}'))
  return changes
 return workbook_transaction(root,c,target,template,mutate)
def snapshot(root,c):
 rows=c.execute('SELECT * FROM files ORDER BY month,team,site,category,source_path').fetchall();items=[]
 for r in rows:
  # Metadata only: no absolute paths, private source bodies or personal opinion text.
  item={k:r[k] for k in ['id','category','month','team','site','site_code','sha256','size','first_seen','last_seen','state']}|{'filename':Path(r['source_path']).name}
  delivered=c.execute('SELECT sha256,size FROM deliveries WHERE file_id=?',(r['id'],)).fetchone()
  if delivered:item.update(dict(delivered))
  items.append(item)
 value={'schema':1,'generated_at':now().isoformat(),'schedule':'02:00','timezone':'Asia/Seoul','collection_state':'NAS_SAVED' if items else 'NOT_RECEIVED','items':items,'review_count':c.execute('SELECT COUNT(*) FROM reviews WHERE resolved=0').fetchone()[0]}
 atomic(root/'dashboard.json',value);return value

def optimize_existing(root,retry_checks=False):
 if not optimization_enabled(root):raise ValueError('IMAGE_OPTIMIZATION_DISABLED')
 c=db(root);stats=collections.Counter();rows=c.execute('SELECT * FROM files ORDER BY category,month,id').fetchall()
 if retry_checks:rows=[r for r in rows if c.execute("SELECT 1 FROM deliveries WHERE file_id=? AND state='PRESERVED_CHECK_REQUIRED'",(r['id'],)).fetchone()]
 try:
  for number,row in enumerate(rows,1):
   dest=Path(row['nas_path']);source=payload(root)/'sources'/row['sha256']/Path(row['source_path']).name
   prior=c.execute('SELECT sha256 FROM deliveries WHERE file_id=?',(row['id'],)).fetchone();expected=prior['sha256'] if prior else row['sha256']
   if not source.is_file() or sha(source)!=row['sha256']:review(c,'IMAGE_SOURCE',row['id'],'축소 전 원본 확인 필요');stats['source_unverified']+=1
   elif not dest.is_file() or sha(dest)!=expected:review(c,'IMAGE_USER_EDIT',row['id'],'NAS 사용자 변경·누락 보존');stats['user_edit_or_missing']+=1
   else:
    saved,proof=deliver(c,root,row['id'],source,dest,row)
    if saved!=dest:raise ValueError('BACKFILL_DESTINATION_CHANGED')
    stats[proof['state']]+=1;stats['source_bytes']+=row['size'];stats['delivered_bytes']+=proof['size']
   c.commit();atomic(root/'optimization-progress.json',{'at':now().isoformat(),'processed':number,'total':len(rows),'counts':dict(stats)})
  dashboard=snapshot(root,c)
  if root==STATE:
   from publish_server import publish as publish_metadata
   server=publish_metadata()
  else:server={'state':'TEST_ONLY'}
  result={'state':'PROCESSED','at':now().isoformat(),'counts':dict(stats),'server':server,'review_count':dashboard['review_count']};atomic(root/'optimization-status.json',result)
  if (root/'status.json').exists():
   status=read(root/'status.json');status['optimization']=optimization_summary(c);status['review_count']=dashboard['review_count'];status['message']=f"NAS 저장·평가자료 준비 완료 · 그림 축소 {status['optimization']['files']}개 · 확인할 항목 {dashboard['review_count']}건";atomic(root/'status.json',status)
  return result
 finally:c.close()

def optimization_summary(c):
 row=c.execute("SELECT COUNT(*) files,COALESCE(SUM(source_size-size),0) saved_bytes FROM deliveries WHERE state='OPTIMIZED'").fetchone();return dict(row)
def publish(root,nas=NAS):
 c=db(root)
 try:
  results=[prepare_evaluation(root,c,nas)];c.commit()
  results.append(prepare_opinions(root,c,nas));c.commit();dashboard=snapshot(root,c)
  server={'state':'TEST_ONLY'}
  if nas==NAS:
   from publish_server import publish as publish_metadata
   server=publish_metadata()
  optimized=optimization_summary(c);extra=f" · 그림 축소 {optimized['files']}개" if optimization_enabled(root) else ''
  result={'state':'PUBLISHED','server':server,'review_count':dashboard['review_count'],'at':now().isoformat(),'message':f"NAS 저장·평가 반영자료 준비 완료{extra} · 확인할 항목 {dashboard['review_count']}건",'workbooks':results,'optimization':optimized}
  # File receipt/publication and searchable metadata are separate states.
  try:
   from search_refresh import request as request_search_refresh
   result['search_refresh']=request_search_refresh(results,root,nas)
  except Exception as error:
   result['search_refresh']={'state':'RETRY_PENDING','error':type(error).__name__,'received_files_preserved':True}
  atomic(root/'status.json',result);return result
 except Exception:
  c.rollback();atomic(root/'status.json',{'state':'NOT_COMPLETE','at':now().isoformat(),'message':'대장 반영 미완료 · 기존 자료를 보존하고 재시도합니다.'});raise
 finally:c.close()

def portable_link(path,root=NAS_LINK_ROOT):
 from urllib.parse import quote
 p=Path(path);below(p,root)
 return quote(str(p.relative_to(root)).replace('\\','/'),safe='/')

def managed_link(address,path,root=NAS_LINK_ROOT):
 from urllib.parse import unquote
 import ntpath
 value=unquote(address or '').replace('/','\\')
 if value.lower().startswith('file:'):value=value[5:].lstrip('\\')
 unc=value.lstrip('\\')
 if unc.casefold().startswith('raidrive-user\\webdav\\'):value='Z:\\'+unc[len('RaiDrive-user\\WebDAV\\'):]
 elif not ntpath.isabs(value):value=ntpath.join(str(root),value)
 return ntpath.normcase(ntpath.normpath(value))==ntpath.normcase(ntpath.normpath(str(path)))

def replace_link(cell,sheet,address):
 # Excel's Address setter fails for long encoded paths. Add handles those
 # paths, but resets the number format and may apply hyperlink font styling.
 number=cell.NumberFormat;text=str(cell.Value2 or '증빙 파일 열기');font={};interior={};alignment={};borders={}
 for name in ['Name','Size','Bold','Italic','Underline','Strikethrough','Subscript','Superscript','OutlineFont','Shadow','ThemeFont','ThemeColor','TintAndShade','Color']:
  try:font[name]=getattr(cell.Font,name)
  except Exception:pass
 for name in ['Pattern','Color','ThemeColor','TintAndShade','PatternColor','PatternThemeColor','PatternTintAndShade']:
  try:interior[name]=getattr(cell.Interior,name)
  except Exception:pass
 for name in ['HorizontalAlignment','VerticalAlignment','WrapText','Orientation','AddIndent','IndentLevel','ShrinkToFit','ReadingOrder','Locked','FormulaHidden']:
  try:alignment[name]=getattr(cell,name)
  except Exception:pass
 for edge in [5,6,7,8,9,10]:
  border=cell.Borders.Item(edge);values={}
  for name in ['Weight','Color','ThemeColor','TintAndShade','LineStyle']:
   try:values[name]=getattr(border,name)
   except Exception:pass
  borders[edge]=values
 cell.Hyperlinks.Item(1).Delete();sheet.Hyperlinks.Add(Anchor=cell,Address=address,TextToDisplay=text);cell.NumberFormat=number
 if font.get('ThemeColor')in range(1,13):font.pop('Color',None)
 else:font.pop('ThemeColor',None)
 for name,value in font.items():
  try:current=getattr(cell.Font,name)
  except Exception:current=None
  if current!=value:setattr(cell.Font,name,value)
 for target,values in [(cell.Interior,interior),(cell,alignment),*[(cell.Borders.Item(edge),values)for edge,values in borders.items()]]:
  if values.get('ThemeColor')in range(1,13):values.pop('Color',None)
  else:values.pop('ThemeColor',None)
  if values.get('PatternThemeColor')in range(1,13):values.pop('PatternColor',None)
  else:values.pop('PatternThemeColor',None)
  for name,value in values.items():
   try:current=getattr(target,name)
   except Exception:current=None
   if current!=value:setattr(target,name,value)

def _rebind_file(c,row,sites):
 if row['site_code'] or row['state']!='SITE_REVIEW':return None
 from site_binding import bind
 decision=bind(row['source_path'],sites)
 if not decision['state'].startswith('EXACT'):return None
 site=next(s for s in sites if s['code']==decision['code'])
 result=c.execute("UPDATE files SET site=?,site_code=?,state='SAVED' WHERE id=? AND site_code='' AND state='SITE_REVIEW'",(site['site'],site['code'],row['id']))
 if result.rowcount!=1:raise ValueError('SITE_BINDING_CHANGED')
 c.execute("UPDATE reviews SET resolved=1 WHERE kind='SITE' AND source_id=? AND resolved=0 AND reason='현장 일치 확인 필요'",(row['id'],))
 return {'file_id':row['id'],'site_code':site['code'],'site':site['site'],'evidence':decision['evidence'],'basis':decision['state'],'month':row['month'],'team_preserved':row['team'],'nas_path_preserved':row['nas_path'],'source_sha256_preserved':row['sha256']}

def reconcile_sites(root,nas=NAS,sites=None):
 c=db(root);sites=catalog(root) if sites is None else sites;changes=[];skipped=[]
 try:
  for row in c.execute("SELECT * FROM files WHERE state='SITE_REVIEW' AND site_code='' ORDER BY id").fetchall():
   from site_binding import bind
   decision=bind(row['source_path'],sites)
   if not decision['state'].startswith('EXACT'):continue
   delivery=c.execute('SELECT sha256,size FROM deliveries WHERE file_id=?',(row['id'],)).fetchone();expected=delivery['sha256']if delivery else row['sha256'];size=delivery['size']if delivery else row['size'];path=Path(row['nas_path'])
   below(path,nas)
   if not path.is_file() or path.stat().st_size!=size or sha(path)!=expected:skipped.append(row['id']);continue
   change=_rebind_file(c,row,sites)
   if change:changes.append(change)
  if changes:book=prepare_evaluation(root,c,nas)
  else:book={'state':'UNCHANGED','changed_cells':0}
  c.commit();dashboard=snapshot(root,c)
  server={'state':'TEST_ONLY'}
  if nas==NAS:
   from publish_server import publish as publish_metadata
   server=publish_metadata()
  result={'state':'SITE_BINDINGS_UPDATED','at':now().isoformat(),'bound_files':len(changes),'changes':changes,'user_changed_or_missing_files_preserved':skipped,'files':len(dashboard['items']),'review_count':dashboard['review_count'],'workbook':book,'server':server,'opinions_or_scores_updated':False}
  try:
   from search_refresh import request as request_search_refresh
   result['search_refresh']=request_search_refresh([book],root,nas)
  except Exception as error:
   result['search_refresh']={'state':'RETRY_PENDING','error_type':type(error).__name__}
  atomic(root/'receipts'/('site-bindings-'+now().strftime('%Y%m%d_%H%M%S%f')+'.json'),result)
  if(root/'status.json').exists():
   status=read(root/'status.json');status['review_count']=dashboard['review_count'];status['site_binding_update']={'at':result['at'],'bound_files':len(changes),'workbook':book};status['message']=f"NAS 저장·평가 반영자료 준비 완료 · 확인할 항목 {dashboard['review_count']}건";atomic(root/'status.json',status)
  return result
 except Exception:c.rollback();raise
 finally:c.close()

def main():
 p=argparse.ArgumentParser();p.add_argument('--root',type=Path,default=STATE);p.add_argument('--result-file',type=Path);s=p.add_subparsers(dest='command',required=True)
 r=s.add_parser('request');r.add_argument('--manual',action='store_true');i=s.add_parser('ingest');i.add_argument('--manifest',required=True,type=Path);s.add_parser('publish');s.add_parser('status');s.add_parser('reconcile-sites');o=s.add_parser('optimize-existing');o.add_argument('--retry-checks',action='store_true')
 a=p.parse_args()
 with locked(a.root,'request.lock' if a.command=='request' else 'worker.lock'):
  if a.command=='request':result=request(a.root,a.manual)
  elif a.command=='ingest':result=ingest(a.root,a.manifest)
  elif a.command=='publish':result=publish(a.root)
  elif a.command=='reconcile-sites':result=reconcile_sites(a.root)
  elif a.command=='optimize-existing':result=optimize_existing(a.root,a.retry_checks)
  else:result=read(a.root/'status.json') if (a.root/'status.json').exists() else {'state':'NOT_RUN'}
 if a.result_file:
  below(a.result_file,a.root/'worker-results');atomic(a.result_file,result)
 print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':
 try:main()
 except Exception as e:
  print(json.dumps({'state':'FAILED','error':type(e).__name__},ensure_ascii=False));sys.exit(1)
