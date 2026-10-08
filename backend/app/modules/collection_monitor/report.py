"""Downloadable report from the same scoped monthly collection payload."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def text(value):
    s=str(value or '')
    return "'"+s if s.startswith(('=','+','-','@')) else s


def make_report(payload):
    wb=Workbook();wb.remove(wb.active)
    rows=payload['sites']
    cells=[c for s in rows for c in s['cells']]
    required=sum(c['expected'] for c in cells);received=sum(c['received'] for c in cells)
    rate=received/required if required and all(c['known'] for c in cells) else '확인 필요'
    scopes={'government':'관급공사','other':'기타 현장'}
    title=scopes[payload.get('scope','government')]+' 문서취합현황'
    head=f"{payload['month']}  기준일 {payload['as_of_date']}"
    freq=lambda value:{'EVENT':'착공·변경 시','ADHOC':'해당 시','MONTHLY':'월간','WEEKLY':'주간','HALF_YEARLY':'반기','DAILY':'일간','QUARTERLY':'분기','YEARLY':'연간'}.get(value,value)
    def sheet(name,headers,widths):
        ws=wb.create_sheet(name);ws.append([title]);ws.append([head]);ws.append(['취합률 = 취합한 필수 제출기간 수 / 전체 필수 제출기간 수. 반려 건은 미취합, 승인 완료율은 별도 표시.'])
        ws.append(headers)
        for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes='C5';ws.sheet_view.showGridLines=False
        ws['A1'].font=Font(name='맑은 고딕',size=18,bold=True)
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers));ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers));ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=len(headers));ws.row_dimensions[1].height=30;ws.row_dimensions[2].height=22;ws.row_dimensions[3].height=34;ws.row_dimensions[4].height=28
        ws['A3'].alignment=Alignment(wrap_text=True,vertical='center')
        for c in ws[4]:c.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='334155');c.alignment=Alignment(wrap_text=True)
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_title_rows='1:4';ws.oddFooter.center.text='&P'
        return ws
    ws=sheet('현장별 요약',['현장코드','현장명','필수 제출기간','취합','미취합','취합률','승인 완료','승인 완료율'],[13,50,17,12,12,15,15,18])
    ws.append(['합계',f'활성 현장 {len(rows)}곳',required,received,required-received,rate,sum(c.get('approved',0) for c in cells),sum(c.get('approved',0) for c in cells)/required if required else '대상 없음'])
    for s in rows:
        cs=s['cells'];n=sum(c['expected'] for c in cs);d=sum(c['received'] for c in cs);a=sum(c.get('approved',0) for c in cs)
        ws.append([text(s['site_code']),text(s['site_name']),n,d,n-d,d/n if n and all(c['known'] for c in cs) else '확인 필요',a,a/n if n else '대상 없음'])
    ds=sheet('서류별 취합률',['서류명','주기','필수 제출기간','취합','미취합','취합률','승인 완료'],[45,16,18,14,14,16,16])
    for definition in payload['documents']:
        cs=[c for c in cells if c['code']==definition['code']];n=sum(c['expected'] for c in cs);d=sum(c['received'] for c in cs)
        if not cs:continue
        ds.append([text(definition['title']),freq(definition['frequency']),n,d,n-d,d/n if n and all(c['known'] for c in cs) else '확인 필요',sum(c.get('approved',0) for c in cs)])
    detail=sheet('현장별 서류현황',['현장코드','현장명','서류명','주기','필수 제출기간','취합','미취합','취합률','승인 완료','반려','미제출 기간','확인 상태'],[13,42,40,14,18,12,12,15,15,12,32,16])
    detail.page_setup.fitToWidth=0;detail.print_title_cols='A:B'
    for s in rows:
        for c in s['cells']:
            n=c['expected'];d=c['received']
            periods=['착공 서류' if p=='once' else p.replace('-W',' ')+'주차' if '-W' in p else p for p in c['missing_periods']]
            detail.append([text(s['site_code']),text(s['site_name']),text(c['title']),freq(c['frequency']),n,d,n-d,d/n if n and c['known'] else '확인 필요',c.get('approved',0),c.get('rejected',0),text(', '.join(periods)),'확인됨' if c['known'] else '확인 필요'])
    for ws in wb:
        ws.auto_filter.ref=f'A4:{get_column_letter(ws.max_column)}{ws.max_row}'
        for row in ws.iter_rows(min_row=5):
            ws.row_dimensions[row[0].row].height=36
            for c in row:c.font=Font(name='맑은 고딕',size=11);c.alignment=Alignment(vertical='center',wrap_text=True)
        for col in ([6,8] if ws.title=='현장별 요약' else [6] if ws.title=='서류별 취합률' else [8]):
            for row in range(5,ws.max_row+1):ws.cell(row,col).number_format='0.0%'
    b=BytesIO();wb.save(b);return b.getvalue()
