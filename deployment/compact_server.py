"""Apply the requested PCM/cost settings through the existing revisioned service."""
import os,sys,json,sqlite3,hashlib,shutil
from pathlib import Path
from datetime import datetime,timezone
ROOT=Path(__file__).parent
DB=Path('/srv/besma/database/besma.db')
OUT=Path('/srv/besma/ops_snapshots/government-compact-20261008')
MODE=sys.argv[1]
def digest(value):return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
def connect(path):
    c=sqlite3.connect(path.resolve().as_uri()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row;return c
def protected(path):
    with connect(path) as c:
        names={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        tables=('users','sites','documents','document_instances','document_upload_histories','document_review_histories','functional_eval_assessments','government_contact_messages')
        return {t:digest([dict(r) for r in c.execute('SELECT * FROM '+t+' ORDER BY id')]) for t in tables if t in names}
if MODE=='test':
    os.umask(0o027);work=ROOT/'regression';work.mkdir(exist_ok=True)
    with connect(DB) as source,sqlite3.connect(work/'test.sqlite3') as target:source.backup(target)
    os.environ.update(SQLITE_PATH=str(work/'test.sqlite3'),STORAGE_ROOT=str(work/'storage'),DOCUMENT_EXPLORER_BASE_DIR=str(work/'forms'))
    contacts=work/'storage/collection-monitor';contacts.mkdir(parents=True,exist_ok=True)
    shutil.copy2('/srv/besma/storage/collection-monitor/site-contacts-20261008.json',contacts/'site-contacts-20261008.json')
sys.path.insert(0,'/srv/besma/backend');os.chdir('/srv/besma/backend')
from app.config.settings import settings
from app.modules.collection_monitor.configuration import read_configuration,save_configuration,CollectionSetting
from app.modules.collection_monitor.core import readonly,scoped_overview
from app.core.database import SessionLocal
from app.modules.users.models import User
def current():
    with readonly(settings.sqlite_path) as c:return read_configuration(c)
def proposal(config):
    result=json.loads(json.dumps(config))
    cost=next(i for i in result['items'] if i['code']=='GOV_SAFETY_COST_MONTHLY')
    cost.update(title='산업안전보건관리비',frequency='MONTHLY',group_id='MONTHLY',priority=True,enabled=True,required=True)
    candidates=[i for i in result['items'] if 'PCM' in i['title'].upper()]
    assert len(candidates)<=1
    if candidates:
        candidates[0].update(title='PCM 자료',frequency='EVENT',group_id='EVENT',priority=True,enabled=True,required=True)
    else:result['items'].append({'code':None,'title':'PCM 자료','frequency':'EVENT','group_id':'EVENT','collection_folder':'PCM 자료','priority':True,'enabled':True,'required':True,'order':104})
    return result
def session():
    db=SessionLocal();actor=db.query(User).filter(User.login_id=='sijung',User.role.in_(['HQ_SAFE','HQ_SAFE_ADMIN','SUPER_ADMIN'])).one_or_none()
    if actor is None:actor=db.query(User).filter(User.role.in_(['SUPER_ADMIN','HQ_SAFE_ADMIN'])).order_by(User.id).first()
    assert actor is not None,'AUTHORIZED_ADMIN_ACTOR_MISSING'
    return db,actor
if MODE=='inspect':
    assert Path(settings.sqlite_path).resolve()==DB
    config=current();request=proposal(config)
    (ROOT/'before-settings.json').write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf8')
    (ROOT/'request.json').write_text(json.dumps(request,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'revision':config['revision'],'existing_items':len(config['items']),'requested_items':len(request['items']),'before_sha256':digest(config)},ensure_ascii=True))
elif MODE=='test':
    assert Path(settings.sqlite_path).resolve()==ROOT/'regression/test.sqlite3'
    before=current();request=json.loads((ROOT/'request.json').read_text());assert digest(before)==digest(json.loads((ROOT/'before-settings.json').read_text()))
    old=protected(settings.sqlite_path)
    with connect(settings.sqlite_path) as c:
        excluded=[dict(r) for r in c.execute('SELECT * FROM document_requirements WHERE site_id IN (69,70) ORDER BY id')]
        ids={r[0] for r in c.execute('SELECT id FROM document_requirements')}
    db,actor=session();saved=save_configuration(db,CollectionSetting.model_validate(request),actor.id)
    assert old==protected(settings.sqlite_path)
    cost=next(i for i in saved['items'] if i['code']=='GOV_SAFETY_COST_MONTHLY');pcm=next(i for i in saved['items'] if i['title']=='PCM 자료');assert cost['priority'] and pcm['priority']
    with connect(settings.sqlite_path) as c:
        assert excluded==[dict(r) for r in c.execute('SELECT * FROM document_requirements WHERE site_id IN (69,70) ORDER BY id')]
        assert ids<={r[0] for r in c.execute('SELECT id FROM document_requirements')}
        assert c.execute('SELECT COUNT(*) FROM document_requirements WHERE code=?',(pcm['code'],)).fetchone()[0]==6
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.core.auth import get_current_user,get_current_user_with_bypass
    from app.modules.collection_monitor.routes import router as monitor
    from app.modules.documents.routes import router as documents
    from app.modules.document_submissions.routes import router as uploads
    app=FastAPI();app.include_router(monitor);app.include_router(documents);app.include_router(uploads);user={'value':actor};app.dependency_overrides[get_current_user]=lambda:user['value'];app.dependency_overrides[get_current_user_with_bypass]=lambda:user['value'];client=TestClient(app)
    overview=client.get('/collection-monitor/overview?month=2026-10').json()
    assert len(overview['sites'])==6 and len(overview['documents'])==9
    assert [d['group_id'] for d in overview['documents']]==['MONTHLY']*3+['WEEKLY']+['EVENT']*5
    user['value']=db.query(User).filter(User.role=='HQ_OTHER').first();assert client.put('/collection-monitor/settings',json=saved).status_code==403
    user['value']=db.query(User).filter(User.role=='SITE',User.site_id==74).first()
    site_config=client.get('/collection-monitor/site-configuration').json();assert {cost['code'],pcm['code']}<={i['code'] for i in site_config['items'] if i['priority']}
    requirements=client.get('/documents/requirements/status?site_id=74&period=all&date=2026-10-08').json()['items']
    for item in (cost,pcm):
        req=next(r for r in requirements if r['document_type_code']==item['code'])
        result=client.post('/document-submissions/upload',data={'site_id':74,'requirement_id':req['requirement_id'],'document_type_code':item['code'],'work_date':'2026-10-08'},files={'file':('fixture.txt',item['title'].encode(),'text/plain')});assert result.status_code==200,result.text[:200]
    user['value']=actor
    plan=client.get('/collection-monitor/nas-export?month=2026-10').json();assert any('/관급/월간/' in r['relative_path'] and '/산업안전보건관리비/' in r['relative_path'] for r in plan['items']);assert any('/관급/착공시/' in r['relative_path'] and '/PCM 자료/' in r['relative_path'] for r in plan['items'])
    from io import BytesIO
    from openpyxl import load_workbook
    report=client.get('/collection-monitor/report.xlsx?month=2026-10');assert report.status_code==200
    wb=load_workbook(BytesIO(report.content));titles={r[0].value for r in wb['서류별 취합률'].iter_rows(min_row=5)};assert {'산업안전보건관리비','PCM 자료'}<=titles
    result={'passed':True,'priority_documents':9,'sites':6,'monthly':3,'weekly':1,'startup':5,'new_requirements':6,'existing_history_preserved':True,'excluded_sites_preserved':True,'site_uploads':2,'report_verified':True,'nas_paths_verified':True,'settings_permissions_verified':True}
    (ROOT/'regression/result.json').write_text(json.dumps(result,indent=2),encoding='utf8')
    for name,value in [('overview.json',overview),('settings.json',saved)]: (ROOT/'regression'/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result));db.close()
elif MODE=='apply':
    assert Path(settings.sqlite_path).resolve()==DB and (ROOT/'regression/result.json').exists()
    config=current();assert digest(config)==digest(json.loads((ROOT/'before-settings.json').read_text())),'CONFIG_CHANGED'
    assert not (OUT/'deployment.json').exists();OUT.mkdir(exist_ok=True,mode=0o750)
    before=protected(DB)
    with connect(DB) as source,sqlite3.connect(OUT/'database-before.sqlite3') as target:source.backup(target)
    (OUT/'database-before.sqlite3').chmod(0o640)
    db,actor=session();saved=save_configuration(db,CollectionSetting.model_validate_json((ROOT/'request.json').read_text()),actor.id)
    assert before==protected(DB)
    receipt={'applied_at':datetime.now(timezone.utc).isoformat(),'revision':saved['revision'],'priority_documents':9,'configuration_documents':len(saved['items']),'protected_data_unchanged':True,'commit':(ROOT/'commit.txt').read_text().strip()}
    for name,value in [('deployment.json',receipt),('settings.json',saved),('protected-before.json',before)]: (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf8');(OUT/name).chmod(0o640)
    shutil.copy2(ROOT/'regression/result.json',OUT/'regression.json');print(json.dumps(receipt));db.close()
elif MODE=='verify':
    assert Path(settings.sqlite_path).resolve()==DB
    config=current();overview=scoped_overview(DB,Path(settings.storage_root)/'collection-monitor','2026-10','government',True)
    assert len(overview['documents'])==9 and len(overview['sites'])==6
    assert [d['group_id'] for d in overview['documents']]==['MONTHLY']*3+['WEEKLY']+['EVENT']*5
    assert protected(DB)==json.loads((OUT/'protected-before.json').read_text())
    for name,value in [('overview.json',overview),('settings.json',config)]: (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf8');(OUT/name).chmod(0o640)
    print(json.dumps({'priority_documents':9,'configuration_documents':len(config['items']),'revision':config['revision'],'protected_data_unchanged':True,'groups':[g['label'] for g in config['groups']]}))
else:raise ValueError(MODE)
