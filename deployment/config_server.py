import os,sys,sqlite3,json,hashlib,shutil,subprocess,time,urllib.request
from pathlib import Path
from datetime import datetime,timezone
R=Path(__file__).parent;RELEASE=R/'release';LIVE=Path('/srv/besma/backend');DB=Path('/srv/besma/database/besma.db');OUT=Path('/srv/besma/ops_snapshots/government-config-20261008')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def summary():
    with sqlite3.connect(DB.as_uri()+'?mode=ro',uri=True) as c:
        tables={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        return {t:{'count':len(rows),'sha256':hashlib.sha256(repr(rows).encode()).hexdigest()} for t in ('users','sites','documents','document_instances','document_upload_histories','document_review_histories','functional_eval_assessments','document_requirements','document_type_masters','submission_cycles','government_contact_messages') if t in tables and (rows:=c.execute('SELECT * FROM '+t+' ORDER BY id').fetchall()) is not None}
def backup(target):
    with sqlite3.connect(DB.as_uri()+'?mode=ro',uri=True) as src,sqlite3.connect(target) as dest:src.backup(dest)
    target.chmod(0o640)
def migrate():
    sys.path.insert(0,str(RELEASE/'backend'));os.chdir(RELEASE/'backend')
    from app.core.database import engine
    from app.modules.collection_monitor.models import GovernmentCollectionConfiguration,GovernmentCollectionConfigurationHistory
    GovernmentCollectionConfiguration.__table__.create(engine,checkfirst=True);GovernmentCollectionConfigurationHistory.__table__.create(engine,checkfirst=True)
mode=sys.argv[1]
if mode=='migrate':
    assert os.geteuid()!=0
    sys.path.insert(0,str(RELEASE/'backend'));os.chdir(RELEASE/'backend')
    from app.config.settings import settings
    assert Path(settings.sqlite_path).resolve()==DB.resolve()
    before=summary();backup(OUT/'database-before.sqlite3');migrate();assert before==summary();(OUT/'protected-before.json').write_text(json.dumps(before),encoding='utf8')
elif mode=='test':
    contacts=R/'regression/storage/collection-monitor';contacts.mkdir(parents=True,exist_ok=True);shutil.copy2(R/'site-contacts-20261008.json',contacts/'site-contacts-20261008.json')
    import operations
    operations.test(RELEASE,R/'regression');migrate()
    from app.config.settings import settings
    assert Path(settings.sqlite_path)==R/'regression/test.sqlite3'
    from app.core.auth import get_current_user,get_current_user_with_bypass
    from app.core.database import SessionLocal
    from app.modules.users.models import User
    from app.modules.collection_monitor.routes import router as monitor
    from app.modules.documents.routes import router as documents
    from app.modules.document_submissions.routes import router as uploads
    from app.modules.collection_monitor.core import readonly
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    session=SessionLocal();hq=session.query(User).filter(User.role=='HQ_SAFE').first();site=session.query(User).filter(User.role=='SITE',User.site_id==74).first();other=session.query(User).filter(User.role=='HQ_OTHER').first();assert all((hq,site,other))
    current={'user':hq};app=FastAPI();app.include_router(monitor);app.include_router(documents);app.include_router(uploads);app.dependency_overrides[get_current_user]=lambda:current['user'];app.dependency_overrides[get_current_user_with_bypass]=lambda:current['user'];client=TestClient(app)
    initial=client.get('/collection-monitor/settings');assert initial.status_code==200;config=initial.json();assert [g['label'] for g in config['groups']]==['월간','주간','착공시'] and len(config['items'])==23
    (R/'regression/settings.json').write_text(json.dumps(config,ensure_ascii=False),encoding='utf8')
    overview=client.get('/collection-monitor/overview',params={'month':'2026-10'}).json();assert [d['group_id'] for d in overview['documents']]==['MONTHLY','MONTHLY','WEEKLY','EVENT','EVENT','EVENT','EVENT']
    clone=sqlite3.connect(settings.sqlite_path);protected_tables=('users','sites','document_instances','document_upload_histories','document_review_histories','functional_eval_assessments');protected={t:clone.execute('SELECT * FROM '+t+' ORDER BY id').fetchall() for t in protected_tables};excluded=clone.execute('SELECT * FROM document_requirements WHERE site_id IN (69,70) ORDER BY id').fetchall();existing_ids=[r[0] for r in clone.execute('SELECT id FROM document_requirements ORDER BY id')]
    # Rename both groups and document title; new items become actual requirements for all six sites.
    config['groups'][0]['label']='월간 취합';config['groups'][1]['label']='주간 점검'
    for item in config['items']:
        if item['code']=='GOV_RISK_MONTHLY':item.update(title='월간 위험성평가',collection_folder='3. 월간 위험성평가')
    config['items'].append({'code':None,'title':'추가 확인서','frequency':'MONTHLY','collection_folder':'추가 확인서','priority':True,'enabled':True,'required':True,'order':50})
    saved=client.put('/collection-monitor/settings',json=config);assert saved.status_code==200,saved.text[:500];new=saved.json();assert new['revision']==1
    code=next(i['code'] for i in new['items'] if i['title']=='추가 확인서');assert code.startswith('GOV_CUSTOM_')
    assert clone.execute('SELECT COUNT(*) FROM document_requirements WHERE code=?',(code,)).fetchone()[0]==6
    assert all(i in [r[0] for r in clone.execute('SELECT id FROM document_requirements')] for i in existing_ids)
    assert excluded==clone.execute('SELECT * FROM document_requirements WHERE site_id IN (69,70) ORDER BY id').fetchall()
    for t,rows in protected.items():assert rows==clone.execute('SELECT * FROM '+t+' ORDER BY id').fetchall(),t
    assert client.put('/collection-monitor/settings',json=config).status_code==409
    invalid=json.loads(json.dumps(new));invalid['items'][0]['title']='../escape';assert client.put('/collection-monitor/settings',json=invalid).status_code==422
    omitted=json.loads(json.dumps(new));omitted['items']=omitted['items'][1:];assert client.put('/collection-monitor/settings',json=omitted).status_code==422
    current['user']=other;assert client.get('/collection-monitor/settings').status_code==403;assert client.put('/collection-monitor/settings',json=new).status_code==403
    current['user']=site;assert client.put('/collection-monitor/settings',json=new).status_code==403;assert client.get('/collection-monitor/settings').status_code==403
    site_config=client.get('/collection-monitor/site-configuration');assert site_config.status_code==200 and any(i['code']==code for i in site_config.json()['items']);assert all('collection_folder' not in i for i in site_config.json()['items'])
    status=client.get('/documents/requirements/status',params={'site_id':74,'period':'all','date':'2026-10-08'});assert status.status_code==200
    requirement=next(i for i in status.json()['items'] if i['document_type_code']==code)
    upload=client.post('/document-submissions/upload',data={'site_id':74,'requirement_id':requirement['requirement_id'],'document_type_code':code,'work_date':'2026-10-08'},files={'file':('new.txt',b'new monthly fixture','text/plain')});assert upload.status_code==200,upload.text[:300]
    current['user']=hq
    plan=client.get('/collection-monitor/nas-export',params={'month':'2026-10'});assert plan.status_code==200;items=plan.json()['items']
    assert any(i['relative_path'].startswith('추가 확인서/관급/월간 취합/') and '/추가 확인서/' in i['relative_path'] for i in items)
    assert any(i['relative_path'].startswith('3. 월간 위험성평가/관급/월간 취합/') and '/월간 위험성평가/' in i['relative_path'] for i in items)
    # Existing content bytes and previous versions remain readable after name changes.
    for item in items:
        response=client.get('/collection-monitor/nas-files/'+item['key'],params={'month':'2026-10'});assert response.status_code==200 and hashlib.sha256(response.content).hexdigest()==item['sha256']
    for item in new['items']:
        if item['code']==code:item.update(title='추가 주간 확인서',frequency='WEEKLY',collection_folder='주간 확인서')
    saved=client.put('/collection-monitor/settings',json=new);assert saved.status_code==200;second=saved.json();assert second['revision']==2
    overview=client.get('/collection-monitor/overview',params={'month':'2026-10'}).json();assert len(overview['documents'])==8 and all(len(s['cells'])==8 for s in overview['sites'])
    assert all(next(c for c in s['cells'] if c['code']==code)['frequency']=='WEEKLY' for s in overview['sites'])
    plan=client.get('/collection-monitor/nas-export',params={'month':'2026-10'}).json();assert any(i['relative_path'].startswith('주간 확인서/관급/주간 점검/') and '/추가 주간 확인서/' in i['relative_path'] for i in plan['items'])
    current['user']=site;status=client.get('/documents/requirements/status',params={'site_id':74,'period':'all','date':'2026-10-08'}).json();assert next(i for i in status['items'] if i['document_type_code']==code)['frequency']=='WEEKLY'
    current['user']=hq;report=client.get('/collection-monitor/report.xlsx',params={'month':'2026-10'});assert report.status_code==200
    from openpyxl import load_workbook
    from io import BytesIO
    workbook=load_workbook(BytesIO(report.content));assert any(row[0].value=='추가 주간 확인서' for row in workbook['서류별 취합률'].iter_rows(min_row=5))
    disabled=json.loads(json.dumps(second));next(i for i in disabled['items'] if i['code']==code)['enabled']=False;assert client.put('/collection-monitor/settings',json=disabled).status_code==200
    current['user']=site;assert all(i['code']!=code for i in client.get('/collection-monitor/site-configuration').json()['items'])
    assert client.post('/document-submissions/upload',data={'site_id':74,'requirement_id':requirement['requirement_id'],'document_type_code':code,'work_date':'2026-10-08'},files={'file':('disabled.txt',b'must not upload','text/plain')}).status_code==403
    assert client.post('/document-submissions/upload',data={'instance_id':upload.json()['instance_id']},files={'file':('disabled.txt',b'must not upload','text/plain')}).status_code==403
    assert clone.execute('SELECT COUNT(*) FROM government_collection_configuration_histories').fetchone()[0]==3
    assert len(list((Path(settings.storage_root)/'collection-monitor/settings-backups').glob('*.sqlite3')))>=3
    result={'existing_regression_groups':7,'default_groups':['월간','주간','착공시'],'default_priority_order_verified':True,'configuration_rename_add_frequency':True,'six_site_requirements_added':True,'site_upload_new_document':True,'nas_group_document_folder_rename':True,'original_version_hash_preserved':True,'report_new_document':True,'role_and_revision_guards':True,'unsafe_folder_and_omission_rejected':True,'disabled_document_upload_denied':True,'history_and_backup':True,'excluded_sites_unchanged':True,'production_db_readonly_clone':True}
    (R/'regression/config-regression.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result));clone.close();session.close()
elif mode=='deploy':
    assert os.geteuid()==0;plan=json.loads((R/'config-manifest.json').read_text());assert (R/'regression/config-regression.json').exists()
    for row in plan['files']:
        path=LIVE/row['path'];assert (sha(path) if path.is_file() else None)==row['before_sha256'],('SOURCE_CHANGED',row['path']);assert sha(RELEASE/'backend'/row['path'])==row['sha256']
    import pwd,grp
    uid=pwd.getpwnam('besma').pw_uid;gid=grp.getgrnam('besma').gr_gid
    assert not (OUT/'deployment.json').exists();OUT.mkdir(mode=0o750,exist_ok=True);os.chown(OUT,uid,gid)
    for row in plan['files']:
        path=LIVE/row['path']
        if path.exists():target=OUT/'before'/row['path'];target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,target)
    subprocess.run(['sudo','-n','-u','besma','/srv/besma/backend/.venv/bin/python','-B',str(R/'config_server.py'),'migrate'],check=True)
    before=summary()
    try:
        for row in plan['files']:
            path=LIVE/row['path'];path.parent.mkdir(parents=True,exist_ok=True);tmp=path.with_name(path.name+'.config-release');shutil.copy2(RELEASE/'backend'/row['path'],tmp);os.chown(tmp,0,gid);os.chmod(tmp,0o640);os.replace(tmp,path)
        subprocess.run(['systemctl','restart','besma-backend.service'],check=True)
        for attempt in range(15):
            try:health=json.loads(urllib.request.urlopen('http://127.0.0.1:8001/health',timeout=4).read());break
            except Exception:
                if attempt==14:raise
                time.sleep(1)
        assert before==summary()
        (OUT/'deployment.json').write_text(json.dumps({**plan,'deployed_at':datetime.now(timezone.utc).isoformat(),'health':health,'protected_data_unchanged':True},indent=2),encoding='utf8');shutil.copy2(R/'regression/config-regression.json',OUT/'regression.json')
        for p in OUT.glob('*.json'):os.chown(p,uid,gid);os.chmod(p,0o640)
        print(json.dumps({'deployed_files':len(plan['files']),'health':health,'protected_data_unchanged':True}))
    except Exception:
        for row in plan['files']:
            path=LIVE/row['path'];original=OUT/'before'/row['path']
            if original.exists():shutil.copy2(original,path)
        subprocess.run(['systemctl','restart','besma-backend.service'],check=True);raise
elif mode=='verify':
    sys.path.insert(0,str(LIVE));os.chdir(LIVE);os.umask(0o027)
    from app.core.auth import get_current_user,get_current_user_with_bypass
    from app.core.database import SessionLocal
    from app.modules.users.models import User
    from app.main import create_app
    from fastapi.testclient import TestClient
    session=SessionLocal();current={'user':session.query(User).filter(User.role=='HQ_SAFE').first()};app=create_app();app.dependency_overrides[get_current_user]=lambda:current['user'];app.dependency_overrides[get_current_user_with_bypass]=lambda:current['user'];client=TestClient(app)
    config=client.get('/collection-monitor/settings');assert config.status_code==200;config=config.json();assert config['revision']==0 and [g['label'] for g in config['groups']]==['월간','주간','착공시']
    overview=client.get('/collection-monitor/overview',params={'month':'2026-10'}).json();assert len(overview['sites'])==6 and [d['group_id'] for d in overview['documents']]==['MONTHLY','MONTHLY','WEEKLY','EVENT','EVENT','EVENT','EVENT']
    assert client.get('/collection-monitor/report.xlsx',params={'month':'2026-10'}).status_code==200
    current['user']=session.query(User).filter(User.role=='HQ_OTHER').first();assert client.get('/collection-monitor/settings').status_code==403;assert client.put('/collection-monitor/settings',json=config).status_code==403
    current['user']=session.query(User).filter(User.role=='SITE',User.site_id==74).first();assert client.get('/collection-monitor/settings').status_code==403;assert client.get('/collection-monitor/site-configuration').status_code==200
    current['user']=session.query(User).filter(User.role=='SITE',User.site_id==69).first();assert client.get('/collection-monitor/site-configuration').status_code==403
    assert summary()==json.loads((OUT/'protected-before.json').read_text())
    receipt=json.loads((OUT/'deployment.json').read_text());assert all(sha(LIVE/r['path'])==r['sha256'] for r in receipt['files'])
    result={'version':overview['version'],'government_sites':6,'default_priority_documents':7,'groups':['월간','주간','착공시'],'configuration_documents':len(config['items']),'configuration_revision':0,'protected_data_unchanged':True,'permissions_verified':True,'commit':receipt['commit'],'verification':'read-only internal app with explicit authentication dependency overrides; write regression in isolated clone only'}
    for name,payload in [('verification.json',result),('overview.json',overview),('settings.json',config)]: (OUT/name).write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result));session.close()
else:raise ValueError(mode)
